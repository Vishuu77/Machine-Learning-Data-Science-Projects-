from docx import Document
import os

doc = Document()

# Title
doc.add_heading('Building Agentic AI Agents: Session Implementation Report', level=0)
doc.add_paragraph('Practical Guide & Code Blueprints Based on Live Session Development')

# Section 1
doc.add_heading('1. Overview of Agents Developed', level=1)
doc.add_paragraph(
    'In this session, we built and tested three distinct functional AI Agents combining Data Science, '
    'Machine Learning, and System Automation:'
)
doc.add_paragraph('• ML Customer Churn / Server Health Agent: Uses Pandas and Scikit-Learn (Random Forest) to ingest structured data and make automated predictive decisions.')
doc.add_paragraph('• Stealth Screen & Audio Recording Agent: Background automation that records desktop video frames (via OpenCV and Pillow) and microphone audio (via sounddevice) in stealth mode, toggled via global hotkeys (Ctrl + Shift + S).')
doc.add_paragraph('• Deep Document Reader & Summarizer Agent: Parses actual PDFs, Word documents (.docx), and text files, extracts their contents, calculates word counts, and generates structured AI summary reports.')

# Section 2
doc.add_heading('2. Step-by-Step Implementation Guide', level=1)
doc.add_paragraph('Step 1: Environment & Dependency Setup')
doc.add_paragraph('Install required Python libraries for data science, GUI, and system automation:\n'
                  'pip install pandas scikit-learn streamlit opencv-python sounddevice soundfile pypdf python-docx keyboard')

doc.add_paragraph('Step 2: Developing the Data Science & ML Decision Agent')
doc.add_paragraph('Load operational data into a Pandas DataFrame, train a Scikit-Learn Random Forest Classifier, and wrap it in an agent loop to evaluate live incoming telemetry.')

doc.add_paragraph('Step 3: Developing the Stealth Screen & Audio Recording Agent')
doc.add_paragraph('Implement multi-threading to record screen frames and microphone audio in the background while hiding the UI window (root.withdraw()) and listening for global keyboard triggers (keyboard.add_hotkey).')

doc.add_paragraph('Step 4: Developing the Deep Document Reader Agent')
doc.add_paragraph('Scan target directories, utilize pypdf and python-docx to extract raw text content from files, generate content-aware text summaries, and export a consolidated report.')

# Section 3
doc.add_heading('3. Exact Prompts Used During Development', level=1)
doc.add_paragraph('These are the exact operational prompts and instructions used during our session to build these agents:')

doc.add_paragraph(
    'Prompt 1 (For Machine Learning Data Agent):\n'
    '"Build a Python script using pandas and scikit-learn that trains a Random Forest Classifier on operational telemetry and outputs automated binary decisions with confidence scores."',
    style='Intense Quote'
)

doc.add_paragraph(
    'Prompt 2 (For Stealth Screen & Audio Recording Agent):\n'
    '"Create a Tkinter desktop application in Python that records screen video and microphone audio in the background, hides the UI during recording, and stops via Ctrl+Shift+S."',
    style='Intense Quote'
)

doc.add_paragraph(
    'Prompt 3 (For Deep Document Reading & Summarizing Agent):\n'
    '"Build a Python agent that scans a selected folder, reads the actual text contents of PDF and Word documents using pypdf and python-docx, and auto-generates a structured content summary report."',
    style='Intense Quote'
)

# Save document
output_path = os.path.expanduser("~/Desktop/Agentic_AI_DataScience_Guide.docx")
doc.save(output_path)
print(f"Updated Word document successfully created at: {output_path}")
